import socket
import io
from fastapi import FastAPI, Request, Form, status
from fastapi.responses import HTMLResponse, RedirectResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.middleware.sessions import SessionMiddleware
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import database

app = FastAPI(title="VigiBot - Red Integrada de Salud San Ignacio")

app.add_middleware(SessionMiddleware, secret_key="vigibot_san_ignacio_secret_key_prod_2026")
app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")

@app.on_event("startup")
def startup_event():
    database.init_db()

# --- RUTA PÚBLICA (CHATBOT) ---
@app.get("/", response_class=HTMLResponse)
async def home(request: Request):
    return templates.TemplateResponse("index.html", {"request": request})

@app.post("/api/guardar-reporte")
async def api_guardar_reporte(data: dict):
    reporte_id = database.guardar_reporte(data)
    return JSONResponse({
        "status": "success",
        "id": reporte_id,
        "mensaje": "Reporte registrado exitosamente para revisión profesional."
    })

# --- RUTAS ADMINISTRATIVAS ---
@app.get("/login", response_class=HTMLResponse)
async def login_page(request: Request):
    if request.session.get("authenticated"):
        return RedirectResponse(url="/dashboard", status_code=status.HTTP_302_FOUND)
    return templates.TemplateResponse("login.html", {"request": request, "error": None})

@app.post("/login", response_class=HTMLResponse)
async def login_post(request: Request, password: str = Form(...)):
    if database.verificar_password(password):
        request.session["authenticated"] = True
        return RedirectResponse(url="/dashboard", status_code=status.HTTP_302_FOUND)
    return templates.TemplateResponse("login.html", {
        "request": request, 
        "error": "Contraseña incorrecta. Acceso restringido a personal autorizado."
    })

@app.get("/logout")
async def logout(request: Request):
    request.session.clear()
    return RedirectResponse(url="/login")

@app.get("/dashboard", response_class=HTMLResponse)
async def dashboard_page(request: Request, msg: str = None, error: str = None):
    if not request.session.get("authenticated"):
        return RedirectResponse(url="/login", status_code=status.HTTP_302_FOUND)
    
    reportes = database.obtener_reportes()
    return templates.TemplateResponse("dashboard.html", {
        "request": request, 
        "reportes": reportes,
        "msg": msg,
        "error": error
    })

@app.get("/exportar-excel")
async def exportar_excel(request: Request):
    if not request.session.get("authenticated"):
        return RedirectResponse(url="/login", status_code=status.HTTP_302_FOUND)
    
    reportes = database.obtener_reportes()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reportes RAM - ESAVI"
    
    # Membrete institucional
    ws.merge_cells("A1:K1")
    ws["A1"] = "RED INTEGRADA DE SALUD SAN IGNACIO - CAJAMARCA"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color="094A73", end_color="094A73", fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    
    ws.merge_cells("A2:K2")
    ws["A2"] = "SISTEMA VIGIBOT | Responsable: Q.F. Valeria Flor Echevarria Aroni - Farmacovigilancia"
    ws["A2"].font = Font(name="Calibri", size=11, bold=True, color="333333")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    
    headers = [
        "ID", "Fecha de Registro", "Paciente / Iniciales", "Edad", "Sexo", 
        "Medicamento Sospechoso", "Lote / Vence", "Reacción Adversa (RAM)", 
        "Gravedad", "Desenlace", "Notificador (Contacto)"
    ]
    
    header_fill = PatternFill(start_color="006699", end_color="006699", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    ws.append([]) # Fila 3 vacía
    ws.append(headers) # Fila 4
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    for r in reportes:
        row_data = [
            r.get("id"),
            r.get("fecha_registro"),
            r.get("paciente_nombre"),
            r.get("paciente_edad"),
            r.get("paciente_sexo"),
            r.get("medicamento"),
            r.get("lote"),
            r.get("reaccion"),
            r.get("gravedad"),
            r.get("desenlace"),
            f"{r.get('notificador_nombre')} ({r.get('notificador_tipo')}) - {r.get('notificador_contacto')}"
        ]
        ws.append(row_data)
        current_row = ws.max_row
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
    
    # Ajustar ancho de columnas automáticamente
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=Reportes_Farmacovigilancia_RIS_San_Ignacio.xlsx"}
    )

@app.post("/cambiar-clave")
async def cambiar_clave(
    request: Request, 
    clave_actual: str = Form(...), 
    clave_nueva: str = Form(...),
    clave_confirmar: str = Form(...)
):
    if not request.session.get("authenticated"):
        return RedirectResponse(url="/login", status_code=status.HTTP_302_FOUND)
    
    if not database.verificar_password(clave_actual):
        return RedirectResponse(url="/dashboard?error=La+clave+actual+es+incorrecta#cambiar-clave", status_code=status.HTTP_302_FOUND)
    
    if len(clave_nueva.strip()) < 6:
        return RedirectResponse(url="/dashboard?error=La+nueva+clave+debe+tener+al+menos+6+caracteres#cambiar-clave", status_code=status.HTTP_302_FOUND)
        
    if clave_nueva != clave_confirmar:
        return RedirectResponse(url="/dashboard?error=Las+nuevas+claves+no+coinciden#cambiar-clave", status_code=status.HTTP_302_FOUND)
    
    database.actualizar_password(clave_nueva.strip())
    return RedirectResponse(url="/dashboard?msg=¡Contraseña+actualizada+con+éxito!#cambiar-clave", status_code=status.HTTP_302_FOUND)

def get_local_ip():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        return "TU_IP_LOCAL"

if __name__ == "__main__":
    import uvicorn
    local_ip = get_local_ip()
    print("=" * 65)
    print("🚀 SISTEMA VIGIBOT ACTIVO PARA TODA LA RED LOCAL")
    print(f"👉 En tu propia PC abre:        http://localhost:8000")
    print(f"👉 Para tu esposa / pacientes:  http://{local_ip}:8000")
    print(f"👉 Panel Privado Q.F.:          http://{local_ip}:8000/login")
    print("=" * 65)
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
